import io
from pathlib import Path
from uuid import uuid4

import pytest
from docx import Document
from pypdf import PdfWriter

from app.documents.embeddings import FakeEmbeddings, validate_vectors
from app.documents.parsing import Page, chunks, parse


def test_chunking_stable_overlap_page_and_heading() -> None:
    did = uuid4()
    pages = [Page("Solar storage " * 150, 1, "Solar"), Page("Second page battery " * 50, 2)]
    first = chunks(pages, did, 600, 100)
    assert first == chunks(pages, did, 600, 100)
    assert {row["page"] for row in first} == {1, 2}
    assert all(0 < len(str(row["text"])) <= 600 for row in first)
    assert first[0]["heading"] == "Solar"
    assert len({row["id"] for row in first}) == len(first)


def test_text_markdown_docx_and_invalid_inputs(tmp_path: Path) -> None:
    md = tmp_path / "notes.md"
    md.write_text("# Battery\nThe battery stores 20 kWh.\n", encoding="utf-8")
    assert parse(md, ".md")[0].heading == "Battery"
    doc = Document()
    doc.add_heading("Energy", level=1)
    doc.add_paragraph("The battery stores 20 kWh.")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Capacity"
    table.cell(0, 1).text = "20 kWh"
    path = tmp_path / "energy.docx"
    doc.save(str(path))
    assert "Capacity | 20 kWh" in " ".join(p.text for p in parse(path, ".docx"))
    md.write_bytes(b"\x00invalid")
    with pytest.raises(ValueError):
        parse(md, ".txt")
    md.write_text("")
    with pytest.raises(ValueError, match="empty_document"):
        parse(md, ".txt")


def test_image_only_pdf_requires_ocr(tmp_path: Path) -> None:
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    path = tmp_path / "blank.pdf"
    buffer = io.BytesIO()
    writer.write(buffer)
    path.write_bytes(buffer.getvalue())
    with pytest.raises(ValueError, match="ocr_required"):
        parse(path, ".pdf")


def test_fake_embedding_is_deterministic_and_dimensions_checked() -> None:
    provider = FakeEmbeddings()
    assert provider.embed(["solar battery"]) == provider.embed(["solar battery"])
    validate_vectors(provider.embed(["solar"]), 1, 64)
    with pytest.raises(ValueError, match="embedding_dimension"):
        validate_vectors([[float("nan")]], 1, 1)
    with pytest.raises(ValueError):
        validate_vectors([[1.0]], 1, 64)
